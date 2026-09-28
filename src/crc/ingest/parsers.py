"""Turn a file or a string into a ``Document`` with recovered section structure.

Design choice: pure-Python parsers, no GROBID. GROBID gives the best TEI
structure but is a Java service behind Docker, which makes the system
undeployable on the target Windows machine and unreproducible for a marker. The
font-size heuristic below recovers the section skeleton we actually need
(abstract / intro / methods / results / conclusion) without that dependency; the
training corpus is LaTeX-derived markdown where headings are already exact, so
only inference-time PDFs rely on the heuristic.
"""
from __future__ import annotations

import re
from pathlib import Path

from .chunkers import sections_from_markdown
from .doctype import _FLOAT_CAPTION, candidate_headings, detect
from .schema import (
    Document,
    DocType,
    ParsedSection,
    Section,
    is_strict_heading_name,
    normalise_heading,
)

_NUMBERED_HEAD = re.compile(r"^(?:\d{1,2}(?:\.\d{1,2})*|[IVXLC]{1,6})[\.\)]?\s+\S")

# IEEE/ACM/Springer templates run the abstract straight on from its label:
#   "Abstract—The given project includes ..."
#   "Index Terms—retrieval, caching"
# Because the label shares a line with the body text, a line-anchored heading
# detector never sees it, and the single highest-weight section in the whole
# document (abstract weight 18.4 vs 3.1 for the next) is silently lost. Break
# the label onto its own line so the normal section machinery picks it up.
_INLINE_LEADIN = re.compile(
    r"^[ \t]*(Abstract|Index Terms|Keywords|Key Words|Summary)\s*[-—–:.]+[ \t]*",
    re.MULTILINE | re.IGNORECASE,
)

_WS = re.compile(r"[ \t]+")
_MULTI_NL = re.compile(r"\n{3,}")
# Hyphen split across a line break: "classi-\nfication" -> "classification".
_DEHYPHEN = re.compile(r"(\w)-\n(\w)")


def _tidy(text: str) -> str:
    text = _DEHYPHEN.sub(r"\1\2", text)
    text = _INLINE_LEADIN.sub(r"\1\n", text)
    text = _WS.sub(" ", text)
    text = _MULTI_NL.sub("\n\n", text)
    return text.strip()


_CANONICAL = frozenset({
    Section.ABSTRACT, Section.INTRODUCTION, Section.BACKGROUND,
    Section.RELATED_WORK, Section.METHODS, Section.RESULTS,
    Section.DISCUSSION, Section.CONCLUSION, Section.REFERENCES,
})


def _n_canonical(sections: list[ParsedSection]) -> int:
    """How many distinct canonical sections a segmentation recovered."""
    return len({s.section for s in sections} & _CANONICAL)


def _sections_from_heading_lines(text: str) -> list[ParsedSection]:
    """Segment plain text by lines that look like headings."""
    heads = candidate_headings(text)
    if not heads:
        return []
    sections: list[ParsedSection] = []
    remaining = text
    # Walk the headings in document order, slicing between successive matches.
    positions: list[tuple[int, str]] = []
    cursor = 0
    for h in heads:
        idx = text.find(h, cursor)
        if idx == -1:
            continue
        positions.append((idx, h))
        cursor = idx + len(h)
    if not positions:
        return []
    # Text before the first heading is the front matter (title/abstract block).
    lead = text[: positions[0][0]].strip()
    if lead:
        sections.append(ParsedSection("", Section.ABSTRACT, lead))
    for i, (idx, head) in enumerate(positions):
        start = idx + len(head)
        end = positions[i + 1][0] if i + 1 < len(positions) else len(text)
        body = text[start:end].strip()
        if body:
            sections.append(ParsedSection(head, normalise_heading(head), body))
    return sections


def parse_pdf(path: str | Path) -> Document:
    """Extract text and sections from a PDF.

    Uses PyMuPDF when available because its span-level font sizes make heading
    detection far more reliable than regex over flattened text: a heading is a
    short line set noticeably larger than the modal body size.
    """
    path = Path(path)
    text = ""
    sections: list[ParsedSection] = []
    parser = "none"
    meta: dict = {}
    pdf_title: str | None = None

    try:
        import fitz  # PyMuPDF

        parser = "pymupdf"
        doc = fitz.open(path)
        meta["n_pages"] = doc.page_count
        spans: list[tuple[float, bool, str]] = []
        first_page: list[tuple[float, str]] = []
        pages_text: list[str] = []
        for pno, page in enumerate(doc):
            pages_text.append(page.get_text("text"))
            d = page.get_text("dict")
            for block in d.get("blocks", []):
                for line in block.get("lines", []):
                    line_spans = line.get("spans", [])
                    line_text = "".join(s.get("text", "") for s in line_spans)
                    if not line_text.strip():
                        continue
                    size = max((s.get("size", 0.0) for s in line_spans), default=0.0)
                    # Bit 4 of `flags` is the bold flag; the font name is a
                    # second signal because some producers do not set it.
                    bold = any(
                        (s.get("flags", 0) & 16)
                        or "bold" in str(s.get("font", "")).lower()
                        or "black" in str(s.get("font", "")).lower()
                        for s in line_spans
                    )
                    spans.append((size, bold, line_text.strip()))
                    if pno == 0:
                        first_page.append((size, line_text.strip()))
        doc.close()
        text = _tidy("\n".join(pages_text))

        if spans:
            sizes = [s for s, _, _ in spans]
            body_size = max(set(sizes), key=sizes.count)   # modal size = body text
            meta["body_font_size"] = round(body_size, 2)

            # A paper is typeset title > headings > body. Lines set at the title
            # size are the title, not sections -- and a title that wraps onto a
            # second line would otherwise be read as a heading, taking its
            # vocabulary with it ("... and\nCAG Models" became a methods
            # section). Exclude the largest band from heading candidates.
            title_size = max(sizes)
            if title_size <= body_size * 1.15:
                title_size = float("inf")   # no distinct title band; keep all
            meta["title_font_size"] = (None if title_size == float("inf")
                                       else round(title_size, 2))
            if title_size != float("inf"):
                pdf_title = _title_from_band(first_page, title_size)

            heads = []
            for size, bold, t in spans:
                if size >= title_size:
                    continue
                # "Abstract—The given project ..." — the label shares its line
                # with the body, so contribute just the label as the heading.
                # `_tidy` has already broken the same line in `text`, so the
                # slicer will find it.
                lead = _INLINE_LEADIN.match(t)
                if lead:
                    heads.append(lead.group(1))
                    continue
                if not (1 <= len(t.split()) <= 9) or t.endswith((".", ",", ";")):
                    continue
                # Reject bare "1" / "2.3": in a typeset paper those are
                # affiliation superscripts and page numbers, which are also set
                # larger than body text and would otherwise slice the document
                # at meaningless points.
                if not re.search(r"[A-Za-z]{3,}", t) or _FLOAT_CAPTION.match(t):
                    continue
                larger = size >= body_size * 1.08
                # Two-column IEEE/ACM templates set section headings in bold at
                # the *same* size as body text, so size alone misses every
                # heading in them. Bold plus a corroborating signal recovers
                # those without letting bold inline emphasis through.
                emphasised = bold and (
                    is_strict_heading_name(t)
                    or _NUMBERED_HEAD.match(t)
                    or t.isupper()
                )
                if larger or emphasised:
                    heads.append(t)
            sections = _sections_from_headings_list(text, heads)
    except ImportError:
        try:
            from pypdf import PdfReader

            parser = "pypdf"
            reader = PdfReader(str(path))
            meta["n_pages"] = len(reader.pages)
            text = _tidy("\n".join((p.extract_text() or "") for p in reader.pages))
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError(
                "No PDF backend available. Install pymupdf (preferred) or pypdf."
            ) from exc

    # Fall back on the text-scanning segmenter whenever the font heuristic
    # recovered too little structure -- not only when it found nothing. A PDF
    # that yields one bogus "section" (an author line, say) is exactly the case
    # that needs the fallback most, and testing `if not sections` skipped it.
    if _n_canonical(sections) < 3:
        fallback = _sections_from_heading_lines(text)
        if _n_canonical(fallback) > _n_canonical(sections):
            sections = fallback

    verdict = detect(text, sections, source_hint="pdf")
    return Document(
        text=text, doc_type=verdict.doc_type, sections=sections,
        title=pdf_title or _guess_title(text), source=str(path), parser=parser,
        meta={**meta, "doctype_confidence": verdict.confidence,
              "doctype_reasons": verdict.reasons},
    )


def _sections_from_headings_list(text: str, heads: list[str]) -> list[ParsedSection]:
    """Slice ``text`` at the given heading strings, in document order."""
    positions: list[tuple[int, str]] = []
    cursor = 0
    seen: set[str] = set()
    for h in heads:
        if h in seen:
            continue
        idx = text.find(h, cursor)
        if idx == -1:
            continue
        positions.append((idx, h))
        seen.add(h)
        cursor = idx + len(h)
    if not positions:
        return []
    sections: list[ParsedSection] = []
    lead = text[: positions[0][0]].strip()
    if lead:
        # Text before the first heading is the abstract only when no explicit
        # abstract heading was recovered; otherwise it is title/author front
        # matter and must not inherit the abstract's large weight.
        has_abstract = any(normalise_heading(h) is Section.ABSTRACT
                           for _, h in positions)
        sections.append(ParsedSection(
            "", Section.OTHER if has_abstract else Section.ABSTRACT, lead))
    for i, (idx, head) in enumerate(positions):
        start = idx + len(head)
        end = positions[i + 1][0] if i + 1 < len(positions) else len(text)
        body = text[start:end].strip()
        if body:
            sections.append(ParsedSection(head, normalise_heading(head), body))
    return sections


def parse_docx(path: str | Path) -> Document:
    """Extract text and sections from a .docx using its heading styles."""
    from docx import Document as DocxDocument   # python-docx

    path = Path(path)
    d = DocxDocument(str(path))
    sections: list[ParsedSection] = []
    cur_head, cur_sec, buf = "", Section.OTHER, []
    all_text: list[str] = []
    title: str | None = None

    def close() -> None:
        body = "\n".join(buf).strip()
        if body:
            sections.append(ParsedSection(cur_head, cur_sec, body))

    for para in d.paragraphs:
        txt = para.text.strip()
        if not txt:
            continue
        all_text.append(txt)
        style = (para.style.name or "").lower() if para.style else ""
        if style.startswith("title") and title is None:
            title = txt
            continue
        if style.startswith("heading"):
            close()
            buf = []
            cur_head, cur_sec = txt, normalise_heading(txt)
        else:
            buf.append(txt)
    close()

    text = _tidy("\n\n".join(all_text))
    verdict = detect(text, sections, source_hint="docx")
    return Document(text=text, doc_type=verdict.doc_type, sections=sections,
                    title=title or _guess_title(text), source=str(path),
                    parser="python-docx",
                    meta={"doctype_confidence": verdict.confidence,
                          "doctype_reasons": verdict.reasons})


def parse_markdown(md: str, source: str = "markdown") -> Document:
    """Parse '#'-headed markdown. This is the training-corpus path."""
    title, sections = sections_from_markdown(md)
    text = _tidy(md)
    verdict = detect(text, sections, source_hint="markdown")
    abstract = next((s.text for s in sections if s.section == Section.ABSTRACT), None)
    return Document(text=text, doc_type=verdict.doc_type, sections=sections,
                    title=title, abstract=abstract, source=source,
                    parser="markdown",
                    meta={"doctype_confidence": verdict.confidence,
                          "doctype_reasons": verdict.reasons})


def parse_text(text: str, title: str | None = None,
               source: str = "text") -> Document:
    """Parse a raw string: pasted abstract, article body, or plain-text paper."""
    clean = _tidy(text)
    sections = _sections_from_heading_lines(clean) if len(clean.split()) > 400 else []
    verdict = detect(clean, sections, source_hint="text")
    return Document(text=clean, doc_type=verdict.doc_type, sections=sections,
                    title=title or (_guess_title(clean)
                                    if verdict.doc_type != DocType.PARAGRAPH else None),
                    source=source, parser="text",
                    meta={"doctype_confidence": verdict.confidence,
                          "doctype_reasons": verdict.reasons})


def _title_from_band(first_page: list[tuple[float, str]], title_size: float) -> str | None:
    """The title of a typeset PDF: the first run of first-page lines at the title size.

    A title that wraps ("PDF Knowledge Extraction System with RAG and" / "CAG
    Models") is several lines in the same large font; the line-based guess kept
    only the first. Every chunk carries the title into Agents 1 and 2, so the
    whole title matters, not just its display. Hyphenated breaks are rejoined.
    Anything implausible (a lone digit, a 60-word banner) returns None, and the
    caller falls back on the line-based guess.
    """
    run: list[str] = []
    for size, t in first_page:
        if size >= title_size - 0.5:
            run.append(t)
        elif run:
            break
    if not run:
        return None
    title = run[0]
    for t in run[1:]:
        title = (title[:-1] + t) if title.endswith("-") and t[:1].islower() else f"{title} {t}"
    title = " ".join(title.split())
    n = len(title.split())
    return title if 2 <= n <= 40 and re.search(r"[A-Za-z]{3,}", title) else None


def _guess_title(text: str) -> str | None:
    """First non-trivial line, if it reads like a title rather than a sentence."""
    for line in text.split("\n"):
        s = line.strip()
        if 3 <= len(s.split()) <= 25 and not s.endswith("."):
            return s
        if s:
            return s[:200] if len(s.split()) <= 30 else None
    return None


__all__ = ["parse_pdf", "parse_docx", "parse_markdown", "parse_text"]
